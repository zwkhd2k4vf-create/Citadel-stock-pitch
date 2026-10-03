"""Build NCLH_model_v6.xlsx from NCLH_model_v5.xlsx.

What changes:
  * Thesis 1 and Thesis 2 now feed the model through two KPIs, each built in its own sheet:
      KPI 1 Price Build   - fare-panel cuts applied to unsold cabins, by brand and region
      KPI 2 Booking Build - booking shortfall from deposits -> occupancy and late discounts
  * Fleet Qtr Build: ship x quarter capacity (bottom-up quarterly capacity growth for 2027)
  * Bridge Build: live consensus-case -> base-case EBITDA bridge, one step per KPI
  * Thesis to KPIs: one-page map from the alternative data to the model and to consensus
  * 'Alt Data Build' becomes 'Fare Panel' (data only); the duplicate 'Alt Data' sheet is removed

usage: python3 build_model_v6.py <v5.xlsx> <out.xlsx> [calibration.json]
"""
import json
import re
import sys

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

src, out = sys.argv[1], sys.argv[2]
cal = json.load(open(sys.argv[3])) if len(sys.argv) > 3 else {}
wb = openpyxl.load_workbook(src)

# ---------------------------------------------------------------- styles
F = 'Calibri'
BLUE, BLACK, GREEN, WHITE = 'FF0000FF', 'FF000000', 'FF008000', 'FFFFFFFF'
GRAY = PatternFill('solid', fgColor='FFF2F2F2')
BAR = PatternFill('solid', fgColor='FF404040')
PCT = '0.0%;\\(0.0%\\)'
PCT2 = '0.00%;\\(0.00%\\)'
PTS = '\\+0.0;\\-0.0;0.0'
MM = '#,##0.0_);\\(#,##0.0\\)'
MM0 = '#,##0_);\\(#,##0\\)'
DOL = '\\$#,##0.00_);\\(\\$#,##0.00\\)'
NUM3 = '#,##0.000'
QTRS6 = ['3Q26E', '4Q26E', '1Q27E', '2Q27E', '3Q27E', '4Q27E']
NYB_COLS = ['L', 'M', 'N', 'O', 'P', 'Q']          # same quarters in the Net Yield Build
KC = ['D', 'E', 'F', 'G', 'H', 'I']                # quarter columns in the new sheets


def put(ws, ref, v, fmt=None, color=BLACK, bold=False, fill=None, wrap=False, size=10, italic=False):
    c = ws[ref]
    c.value = v
    c.font = Font(name=F, size=size, bold=bold, color=color, italic=italic)
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    if wrap:
        c.alignment = Alignment(wrap_text=True, vertical='top')
    return c


def bar(ws, ref, text, width_to='J'):
    put(ws, ref, text, color=WHITE, bold=True, fill=BAR)
    row = ws[ref].row
    for col in range(ws[ref].column + 1, openpyxl.utils.column_index_from_string(width_to) + 1):
        ws.cell(row, col).fill = BAR


def sec(ws, ref, text):
    put(ws, ref, text, bold=True)


def note(ws, ref, text):
    put(ws, ref, text, color='FF595959', italic=True, size=9)


def inp(ws, ref, v, fmt=PCT, assumption=True):
    return put(ws, ref, v, fmt, color=BLUE, fill=GRAY if assumption else None)


def fx(ws, ref, formula, fmt=PCT, bold=False):
    """Formula cell: green when it is a plain link to another sheet, black otherwise."""
    link = re.fullmatch(r"=('[^']+'|[A-Za-z][\w ]*)!\$?[A-Z]+\$?\d+", formula) is not None
    return put(ws, ref, formula, fmt, color=GREEN if link else BLACK, bold=bold)


def new_sheet(title, widths, index=None):
    ws = wb.create_sheet(title, index)
    ws.sheet_view.showGridLines = False
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    return ws


def qhead(ws, row, cols=KC, labels=QTRS6, extra=None):
    for c, q in zip(cols, labels):
        put(ws, f'{c}{row}', q, bold=True).alignment = Alignment(horizontal='right')
    if extra:
        for c, q in extra.items():
            put(ws, f'{c}{row}', q, bold=True).alignment = Alignment(horizontal='right')


# ---------------------------------------------------------------- 1. Fare Panel (was Alt Data Build)
fp = wb['Alt Data Build']
fp.title = 'Fare Panel'
policy = [[fp.cell(r, c).value for c in (2, 4, 6)] for r in range(381, 388)]
for row in fp.iter_rows(min_row=262, max_row=fp.max_row):
    for c in row:
        c.value = None
        c.fill = PatternFill()
        c.font = Font(name=F, size=10)
put(fp, 'B2', 'Fare Panel: forward fares (allaboarddeals.com, collected Oct 2, 2026), product tests, line policies and peer deposits',
    color=WHITE, bold=True, fill=BAR)
put(fp, 'B4', 'Sections A-G are the raw fare-panel tables, H summarizes the statistics the memo quotes, N-O are line policies and peer '
    'deposits. KPI 1 Price Build and KPI 2 Booking Build turn these into model inputs.')
FPR = "'Fare Panel'!"
A_B, A_C, A_D, A_E, A_F = (f"{FPR}${c}$8:${c}$114" for c in 'BCDEF')
A_I, A_J = f"{FPR}$I$8:$I$114", f"{FPR}$J$8:$J$114"

sec(fp, 'B262', 'H. Summary statistics used in the memo (formulas over sections A-E)')
put(fp, 'B263', 'Statistic', bold=True)
put(fp, 'F263', 'Value', bold=True)
rl = '=SUMPRODUCT(($B$8:$B$114="{l}")*($C$8:$C$114="Caribbean/Bahamas")*$E$8:$E$114*$I$8:$I$114)/' \
     'SUMPRODUCT(($B$8:$B$114="{l}")*($C$8:$C$114="Caribbean/Bahamas")*$E$8:$E$114)'
stats = [
    ('Caribbean itineraries at their lowest fare on record, all windows: NCL', rl.format(l='Norwegian'), PCT),
    ('Same: Royal Caribbean', rl.format(l='Royal Caribbean'), PCT),
    ('Same: Carnival', rl.format(l='Carnival'), PCT),
    ('4Q26 Caribbean sailings of 7+ nights cut 20% or more: NCL', '=G158', PCT),
    ('Same: Royal Caribbean', '=G161', PCT),
    ('Same: Carnival', '=G164', PCT),
    ('NCL 4Q26 Caribbean 7+ nights: median cut', '=F158', PCT),
    ('Newest ships, Caribbean 4Q26-1H27 sailings, median cut: NCL Aqua and Luna (11 itineraries)', -0.22, PCT),
    ('Same: Royal Star and Legend of the Seas (12 itineraries)', -0.01, PCT),
]
for i, (lab, v, fmt) in enumerate(stats):
    r = 264 + i
    put(fp, f'B{r}', lab)
    if isinstance(v, str):
        fx(fp, f'F{r}', v, fmt)
    else:
        inp(fp, f'F{r}', v, fmt, assumption=False)
sec(fp, 'B274', 'N. What each line gives a guest whose fare later drops (line policy summary, Apr 2026; ncl.com terms)')
for i, row in enumerate(policy):
    for j, v in enumerate(row):
        put(fp, f'{"BDF"[j]}{275 + i}', v, bold=(i == 0))
sec(fp, 'B284', 'O. Customer deposits y/y: NCLH vs. Royal Caribbean vs. Carnival (Deposits sheet)')
for j, h in enumerate(['Quarter', 'NCLH', 'Royal Caribbean', 'Carnival']):
    put(fp, f'{"BCDE"[j]}285', h, bold=True)
for i in range(9):
    r, d = 286 + i, 48 + i
    fx(fp, f'B{r}', f'=Deposits!B{d}', 'yyyy-mm-dd')
    fx(fp, f'C{r}', f'=Deposits!D{d}', PCT)
    fx(fp, f'D{r}', f'=Deposits!F{d}', PCT)
    fx(fp, f'E{r}', f'=Deposits!I{d}', PCT)
note(fp, 'B296', 'Source: allaboarddeals.com public pages (robots.txt permits), collected Oct 2, 2026; Internet Archive (2025 statistics, '
     'market-pulse history); line policy summary (Apr 2026); ncl.com Best Price Guarantee terms; SEC XBRL. Blue = data.')

# ---------------------------------------------------------------- 2. Fleet Qtr Build
fq = new_sheet('Fleet Qtr Build', {'A': 2, 'B': 30, 'C': 12, 'D': 10, 'E': 10, 'F': 10, 'G': 10, 'H': 10, 'I': 10,
                                    'J': 12, 'K': 12, 'L': 10, 'M': 10, 'N': 10, 'O': 10})
bar(fq, 'B2', 'Fleet Qtr Build: capacity days by ship and quarter (bottom-up quarterly capacity growth)', 'O')
note(fq, 'B3', 'One row per ship per quarter (rows 21-524). Capacity days = berths x years in service during the quarter x capacity days per '
     'berth-year (Fleet Build K58). Delivery and exit dates come from Fleet Build. 2027 y/y growth (row 12) feeds the Net Yield Build; '
     '3Q-4Q26 stay on company guidance.')
QL = [f'{q}Q{y}' for y in (25, 26, 27) for q in (1, 2, 3, 4)]
SC = [get_column_letter(4 + i) for i in range(12)]          # D..O
put(fq, 'B5', 'Quarter', bold=True)
for c, q in zip(SC, QL):
    put(fq, f'{c}5', q, bold=True).alignment = Alignment(horizontal='right')
labels = ['Quarter start (year)', 'Capacity days, total (M)', '  Norwegian (M)', '  Oceania (M)', '  Regent (M)',
          'Luxury (Oceania + Regent) share of capacity', 'Capacity growth y/y', 'Reported capacity days (M), for reference']
for i, lab in enumerate(labels):
    put(fq, f'B{6 + i}', lab, bold=(i in (1, 6)))
N0, N1 = 21, 21 + 42 * 12 - 1
rng = lambda col: f'${col}${N0}:${col}${N1}'
for i, (c, q) in enumerate(zip(SC, QL)):
    yr, qn = 2000 + int(q[2:]), int(q[0])
    inp(fq, f'{c}6', yr + (qn - 1) / 4, '0.00', assumption=False)
    fx(fq, f'{c}7', f'=SUMIFS({rng("K")},{rng("D")},{c}$5)', NUM3, bold=True)
    for k, brand in enumerate(['Norwegian', 'Oceania', 'Regent']):
        fx(fq, f'{c}{8 + k}', f'=SUMIFS({rng("K")},{rng("D")},{c}$5,{rng("C")},"{brand}")', NUM3)
    fx(fq, f'{c}11', f'=({c}9+{c}10)/{c}7', PCT)
    if i >= 4:
        fx(fq, f'{c}12', f'={c}7/{SC[i - 4]}7-1', PCT, bold=True)
qa_cols = {'1Q25': 'E', '2Q25': 'F', '3Q25': 'G', '4Q25': 'H', '1Q26': 'I', '2Q26': 'J'}
for c, q in zip(SC, QL):
    if q in qa_cols:
        fx(fq, f'{c}13', f"='Quarterly (A)'!{qa_cols[q]}29", NUM3)
put(fq, 'B15', 'Fleet-implied 2027 capacity days (M) vs. Net Yield Build', bold=True)
fx(fq, 'D15', '=SUM(L7:O7)', NUM3)
fx(fq, 'E15', "='Net Yield Build'!U94", NUM3)
note(fq, 'F15', 'The quarterly build uses the 2026-27 days-per-berth-year, so only y/y growth (not levels) feeds the model.')
for j, h in enumerate(['Ship', 'Brand', 'Quarter', 'Q start', 'Q end', 'Berths', 'In service', 'Out of service',
                       'Years in service in quarter', 'Capacity days (M)']):
    put(fq, f'{"BCDEFGHIJK"[j]}20', h, bold=True, wrap=True)
r = N0
for s in range(42):
    fr = 6 + s
    for q, (c, ql) in enumerate(zip(SC, QL)):
        fx(fq, f'B{r}', f"='Fleet Build'!B{fr}", 'General')
        fx(fq, f'C{r}', f"='Fleet Build'!C{fr}", 'General')
        put(fq, f'D{r}', ql)
        fx(fq, f'E{r}', f'={c}$6', '0.00')
        fx(fq, f'F{r}', f'=E{r}+0.25', '0.00')
        fx(fq, f'G{r}', f"='Fleet Build'!E{fr}", '#,##0')
        fx(fq, f'H{r}', f"='Fleet Build'!F{fr}", '0.00')
        fx(fq, f'I{r}', f"='Fleet Build'!G{fr}", '0.00')
        fx(fq, f'J{r}', f'=MAX(0,MIN(F{r},I{r})-MAX(E{r},H{r}))', '0.000')
        fx(fq, f'K{r}', f"=G{r}*J{r}*'Fleet Build'!$K$58/1000000", '0.0000')
        r += 1
fq.freeze_panes = 'C6'

# ---------------------------------------------------------------- 3. KPI 1 Price Build (Thesis 1)
k1 = new_sheet('KPI 1 Price Build', {'A': 2, 'B': 62, 'C': 48, 'D': 11, 'E': 11, 'F': 11, 'G': 11, 'H': 11, 'I': 11, 'J': 11})
K1 = "'KPI 1 Price Build'!"
K2 = "'KPI 2 Booking Build'!"
bar(k1, 'B2', 'KPI 1 Price Build (Thesis 1): fare cuts on unsold cabins -> ticket revenue per passenger day')
note(k1, 'B3', 'Thesis 1: NCL is cutting prices to fill Caribbean ships because it is losing share there. The fare panel measures the '
     'cuts; this sheet applies them only to cabins NCL still has to sell, by brand and region. Output (section E) feeds the Net Yield Build.')
put(k1, 'C5', 'Logic / source', bold=True)
qhead(k1, 5)
put(k1, 'B6', 'Fare-panel sailing window')
for c, w in zip(KC, ['sailed', 'Q4 2026', '1H 2027', '1H 2027', '2H 2027', '2H 2027']):
    put(k1, f'{c}6', w).alignment = Alignment(horizontal='right')
put(k1, 'B7', 'Months from Oct 2, 2026 to mid-quarter sailing')
for c, v in zip(KC, [0, 1.5, 4.5, 7.5, 10.5, 13.5]):
    inp(k1, f'{c}7', v, '0.0', assumption=False)

sec(k1, 'B9', 'A. Where NCLH capacity sails')
put(k1, 'B10', 'NCLH Caribbean share of capacity (deployment estimate)')
put(k1, 'C10', 'Team estimate from published itinerary schedules')
for c, v in zip(KC, [0.18, 0.5, 0.62, 0.38, 0.2, 0.48]):
    inp(k1, f'{c}10', v)
put(k1, 'B11', 'Oceania + Regent share of NCLH capacity')
put(k1, 'C11', 'Fleet Qtr Build row 11')
for c, fc in zip(KC, ['J', 'K', 'L', 'M', 'N', 'O']):
    fx(k1, f'{c}11', f"='Fleet Qtr Build'!{fc}11")
put(k1, 'B12', "Oceania + Regent: Caribbean share of their itineraries")
put(k1, 'C12', 'Fare Panel A: itinerary counts in the window')
lux = lambda col, extra: (f'SUMIFS({col},{A_B},"Oceania"{extra})+SUMIFS({col},{A_B},"Regent Seven Seas"{extra})')
for c in KC:
    win = f'{c}$6'
    num = lux(A_E, f',{A_C},"Caribbean/Bahamas",{A_D},{win}')
    den = lux(A_E, f',{A_D},{win}')
    fx(k1, f'{c}12', f'=IF({win}="sailed",0,({num})/({den}))')
put(k1, 'B13', 'NCL brand: Caribbean share of its capacity', bold=True)
put(k1, 'C13', '(NCLH share - luxury share x luxury Caribbean share) / NCL share')
for c in KC:
    fx(k1, f'{c}13', f'=MIN(1,({c}10-{c}11*{c}12)/(1-{c}11))', bold=True)

sec(k1, 'B15', "B. Observed fare change: median vs. each itinerary's own 90-day average (Oct 2, 2026)")
rows_b = [
    (16, 'NCL: Caribbean', 'Fare Panel A', lambda w: f'SUMIFS({A_F},{A_B},"Norwegian",{A_C},"Caribbean/Bahamas",{A_D},{w})'),
    (17, 'NCL: Europe, Alaska and other (itinerary-weighted)', 'Fare Panel A',
     lambda w: f'SUMPRODUCT(({A_B}="Norwegian")*({A_C}<>"Caribbean/Bahamas")*({A_D}={w})*{A_E}*{A_F})/'
               f'SUMPRODUCT(({A_B}="Norwegian")*({A_C}<>"Caribbean/Bahamas")*({A_D}={w})*{A_E})'),
    (18, 'Oceania and Regent: all regions (itinerary-weighted)', 'Fare Panel A',
     lambda w: f'SUMPRODUCT((({A_B}="Oceania")+({A_B}="Regent Seven Seas"))*({A_D}={w})*{A_E}*{A_F})/'
               f'SUMPRODUCT((({A_B}="Oceania")+({A_B}="Regent Seven Seas"))*({A_D}={w})*{A_E})'),
    (19, 'Memo: Royal Caribbean and Carnival, Caribbean (average)', 'Fare Panel A',
     lambda w: f'AVERAGE(SUMIFS({A_F},{A_B},"Royal Caribbean",{A_C},"Caribbean/Bahamas",{A_D},{w}),'
               f'SUMIFS({A_F},{A_B},"Carnival",{A_C},"Caribbean/Bahamas",{A_D},{w}))'),
]
for r, lab, srcs, f in rows_b:
    put(k1, f'B{r}', lab)
    put(k1, f'C{r}', srcs)
    for c in KC:
        fx(k1, f'{c}{r}', f'=IF({c}$6="sailed",0,{f(c + "$6")})')
put(k1, 'B20', 'Memo: NCL Caribbean cut beyond its peers')
for c in KC:
    fx(k1, f'{c}20', f'={c}16-{c}19')

sec(k1, 'B22', 'C. Cabins still unsold at today\'s fares')
put(k1, 'B23', 'Booking curve: months before sailing')
for c, v in zip(['D', 'E', 'F', 'G', 'H'], ['0-3', '3-6', '6-9', '9-12', '12-15']):
    put(k1, f'{c}23', v).alignment = Alignment(horizontal='right')
put(k1, 'B24', 'Typical share of cabins already sold (industry curve)')
put(k1, 'C24', 'Industry booking curve (team, from line commentary)')
for c, v in zip(['D', 'E', 'F', 'G', 'H'], [0.85, 0.6, 0.5, 0.35, 0.25]):
    inp(k1, f'{c}24', v)
put(k1, 'B25', "Typical share sold at this quarter's lead time")
for c in KC:
    fx(k1, f'{c}25', f'=IF({c}$6="sailed",1,IF({c}7<=3,$D$24,IF({c}7<=6,$E$24,IF({c}7<=9,$F$24,IF({c}7<=12,$G$24,$H$24)))))')
put(k1, 'B26', 'Booking shortfall in windows the panel already prices (KPI 2 build)')
put(k1, 'C26', 'Q4 2026 and 1H27 only; 2H27 shortfall is Thesis 2')
for c in KC:
    fx(k1, f'{c}26', f'=IF(OR({c}$6="Q4 2026",{c}$6="1H 2027"),{K2}$D$11,0)')
put(k1, 'B27', "Share of cabins unsold at today's fares: base case", bold=True)
put(k1, 'C27', '1 - typical share sold + shortfall')
for c in KC:
    fx(k1, f'{c}27', f'=IF({c}$6="sailed",0,1-{c}25+{c}26)', bold=True)
put(k1, 'B28', 'Same: consensus case')
put(k1, 'C28', 'Calibrated so the consensus case reproduces consensus 2027 EPS')
u_c = cal.get('u_cons', 0.25)
for c, v in zip(KC, [0, 0.2, u_c, u_c, u_c, u_c]):
    inp(k1, f'{c}28', v)
put(k1, 'B29', 'Same: bear case')
put(k1, 'C29', 'Bear shortfall from KPI 2 build')
for c in KC:
    fx(k1, f'{c}29', f'=IF({c}$6="sailed",0,1-{c}25+IF(OR({c}$6="Q4 2026",{c}$6="1H 2027"),{K2}$D$12,0))')
put(k1, 'B30', 'Bear: share of the 1H27 Caribbean cut that carries into 2H27 sailings')
inp(k1, 'D30', 0.5)

sec(k1, 'B32', 'D. Weighting by brand (Oceania and Regent show no cuts)')
put(k1, 'B33', 'NCL brand share of NCLH ticket revenue', bold=True)
put(k1, 'C33', 'Calibrated to reported 2025 revenue at 66% (section H); we use 65%')
inp(k1, 'D33', 0.65)
put(k1, 'B34', 'Cross-check: NCL brand share of 2027 berths (ceiling)')
fx(k1, 'D34', "='Fleet Build'!K50/'Fleet Build'!K53")
put(k1, 'B35', 'Cross-check: median lowest fare per night (NCL / Oceania / Regent)')
for c, l in zip(['D', 'E', 'F'], ['Norwegian', 'Oceania', 'Regent Seven Seas']):
    fx(k1, f'{c}35', f'=SUMPRODUCT(({A_B}="{l}")*{A_E}*{A_J})/SUMIFS({A_E},{A_B},"{l}")', '$#,##0')
put(k1, 'B36', 'Cross-check: NCL share of ticket revenue implied by berths x panel fares')
put(k1, 'C36', 'Lowest fares overstate luxury weight (narrower cabin spread)')
fx(k1, 'D36', "='Fleet Build'!K50*D35/('Fleet Build'!K50*D35+'Fleet Build'!K51*E35+'Fleet Build'!K52*F35)")

sec(k1, 'B38', 'E. KPI 1 output: Thesis 1 effect on NCLH ticket revenue per passenger day (y/y)')
t1 = lambda c, u, carib: (f'=IF({c}$6="sailed",0,$D$33*{c}{u}*({c}13*{carib}+(1-{c}13)*{c}17)+(1-$D$33)*{c}{u}*{c}18)')
put(k1, 'B39', 'Base case', bold=True)
put(k1, 'B40', 'Consensus case')
put(k1, 'B41', 'Bear case')
put(k1, 'C41', '2H27: part of the 1H27 Caribbean cut carries over (row 30)')
put(k1, 'B42', 'Base less consensus (pts)', bold=True)
put(k1, 'B43', 'Memo: part of the base effect that is NCL-specific (Caribbean cut beyond peers)')
for c in KC:
    fx(k1, f'{c}39', t1(c, 27, f'{c}16'), PCT2, bold=True)
    fx(k1, f'{c}40', t1(c, 28, f'{c}16'), PCT2)
    fx(k1, f'{c}41', t1(c, 29, f'IF({c}$6="2H 2027",MIN({c}16,$D$30*$F$16),{c}16)'), PCT2)
    fx(k1, f'{c}42', f'={c}39-{c}40', PCT2, bold=True)
    fx(k1, f'{c}43', f'=IF({c}$6="sailed",0,$D$33*{c}27*{c}13*{c}20)', PCT2)

sec(k1, 'B45', 'F. Other ticket-price items, the same in every case')
put(k1, 'B46', 'Full-fare commissions: run-rate cost (pts of ticket revenue)')
put(k1, 'C46', 'Non-commissionable fares ended for bookings from Dec 26, 2025')
inp(k1, 'D46', 0.015)
put(k1, 'B47', "Share of the quarter's sailings booked under the new terms")
for c, v in zip(KC, [0.4, 0.6, 0.8667, 0.8833, 0.9, 0.9333]):
    inp(k1, f'{c}47', v)
put(k1, 'B48', 'Same quarter a year earlier')
for c, v in zip(KC, [0, 0, 0, 0.15, '=D47', '=E47']):
    if isinstance(v, str):
        fx(k1, f'{c}48', v)
    else:
        inp(k1, f'{c}48', v)
put(k1, 'B49', 'Commission step-up y/y (pts of ticket)', bold=True)
for c in KC:
    fx(k1, f'{c}49', f'=$D$46*({c}47-{c}48)', PCT2, bold=True)

sec(k1, 'B51', 'G. Sensitivity: 1Q27 Thesis 1 effect (base) vs. NCL revenue weight (across) and 1H27 Caribbean cut (down)')
weights = [0.55, 0.6, 0.65, 0.7, 0.8]
for c, w in zip(['D', 'E', 'F', 'G', 'H'], weights):
    inp(k1, f'{c}52', w, assumption=False)
for i, cut in enumerate([-0.05, -0.1, -0.145, -0.2]):
    r = 53 + i
    inp(k1, f'C{r}', cut, assumption=False)
    for c in ['D', 'E', 'F', 'G', 'H']:
        fx(k1, f'{c}{r}', f'={c}$52*$F$27*($F$13*$C{r}+(1-$F$13)*$F$17)+(1-{c}$52)*$F$27*$F$18', PCT2)
sec(k1, 'B60', 'H. Calibrating the NCL brand\'s share of ticket revenue to reported revenue (2025)')
put(k1, 'B61', 'Reported ticket revenue per passenger day, 2025 ($)')
fx(k1, 'D61', "='Operating Model'!G36/'Operating Model'!G30", DOL)
put(k1, 'B62', 'Share of 2025 berths: NCL / Oceania / Regent')
for c, r in zip('DEF', (50, 51, 52)):
    fx(k1, f'{c}62', f"='Fleet Build'!I{r}/'Fleet Build'!I53")
put(k1, 'B63', 'Median lowest fare per night: NCL / Oceania / Regent ($)')
for c in 'DEF':
    fx(k1, f'{c}63', f'={c}35', '$#,##0')
put(k1, 'B64', 'Ticket revenue per passenger day from Oceania and Regent, if they sell near their lowest fare ($)')
put(k1, 'C64', 'All-suite ships: narrow range of cabin prices')
fx(k1, 'D64', '=E62*E63+F62*F63', DOL)
put(k1, 'B65', 'NCL brand share of NCLH ticket revenue, calibrated', bold=True)
put(k1, 'C65', '1 - luxury revenue / reported total')
fx(k1, 'D65', '=1-D64/D61', PCT, bold=True)
put(k1, 'B66', 'Implied NCL realized fare / its lowest fare')
fx(k1, 'D66', '=(D61-D64)/(D62*D63)', '0.00x')
note(k1, 'B67', 'The calibration does not use NCL\'s own (currently cut) fares. Passenger days are taken in proportion to berths; NCL ships carry '
     'more third and fourth guests, which would push the NCL share up. Bounds: 58% if NCL also sold at its lowest fare (row 36), 83% berth share (row 34).')
note(k1, 'B58', 'Blue on gray = assumptions; blue = data; black = formulas; green = links. Fare data: Fare Panel (allaboarddeals.com, '
     'Oct 2, 2026). Capacity: Fleet Build / Fleet Qtr Build.')

# ---------------------------------------------------------------- 4. KPI 2 Booking Build (Thesis 2)
k2 = new_sheet('KPI 2 Booking Build', {'A': 2, 'B': 62, 'C': 48, 'D': 11, 'E': 11, 'F': 11, 'G': 11, 'H': 11, 'I': 11, 'J': 11})
bar(k2, 'B2', 'KPI 2 Booking Build (Thesis 2): booking shortfall -> occupancy and late discounts')
note(k2, 'B3', "Thesis 2: NCL's Best Price Guarantee removes the reason to book early, so the booking shortfall persists and NCL fills it "
     'late at a discount, leaving some cabins empty. Deposits size the shortfall; this sheet splits it into empty cabins (KPI 2: '
     'occupancy) and late-sold cabins (price, 2H27).')
put(k2, 'C5', 'Logic / source', bold=True)
sec(k2, 'B5', 'A. How short is the booked position? (Jun 30, 2026)')
put(k2, 'B6', 'Deposits per future berth-day, y/y')
put(k2, 'C6', 'Deposits sheet: advance ticket sales / next-4Q capacity days')
fx(k2, 'D6', '=Deposits!F26')
put(k2, 'B7', 'Ticket revenue per passenger day, 1H26 y/y (realized price)')
put(k2, 'C7', 'Quarterly (A)')
fx(k2, 'D7', "=(('Quarterly (A)'!I5+'Quarterly (A)'!J5)/('Quarterly (A)'!I30+'Quarterly (A)'!J30))/"
             "(('Quarterly (A)'!E5+'Quarterly (A)'!F5)/('Quarterly (A)'!E30+'Quarterly (A)'!F30))-1")
put(k2, 'B8', 'Shortfall if booked fares fell as much as realized fares (share of cabins)')
fx(k2, 'D8', '=D7-D6')
put(k2, 'B9', 'Shortfall if booked fares were flat')
fx(k2, 'D9', '=-D6')
put(k2, 'B10', 'Midpoint')
fx(k2, 'D10', '=AVERAGE(D8:D9)')
put(k2, 'B11', 'Base-case booking shortfall', bold=True)
put(k2, 'C11', 'Midpoint, rounded down; 4Q26 through 4Q27')
inp(k2, 'D11', 0.05)
put(k2, 'B12', 'Bear-case booking shortfall')
put(k2, 'C12', 'About the fares-flat reading (row 9)')
inp(k2, 'D12', 0.08)
put(k2, 'B13', 'Consensus case: no shortfall (base-loading works)')
inp(k2, 'D13', 0)
note(k2, 'B14', 'Management: NCL is "below its optimal booked position for the next 12 months" (2Q26 call).')

sec(k2, 'B16', 'B. How much of the shortfall stays empty (KPI 2: occupancy)')
put(k2, 'B17', 'Cross-check: 2Q26 occupancy change y/y (pts)')
fx(k2, 'D17', "=('Quarterly (A)'!J31-'Quarterly (A)'!F31)*100", PTS)
put(k2, 'B18', 'Deposit indicator a quarter earlier (Mar 31, 2026)')
fx(k2, 'D18', '=Deposits!F25')
put(k2, 'B19', 'Occupancy points per 1% of indicator decline (2Q26)')
fx(k2, 'D19', '=D17/(D18*100)', '0.00')
put(k2, 'B20', 'Same, 3Q26 guide vs. the Jun 30 indicator')
fx(k2, 'D20', "=(('Net Yield Build'!L47-'Quarterly (A)'!G31)*100)/(Deposits!F26*100)", '0.00')
put(k2, 'B21', "Occupancy drop that relationship implies at today's indicator (pts)")
fx(k2, 'D21', '=AVERAGE(D19:D20)*Deposits!F26*100', PTS)
put(k2, 'B22', 'Share of the shortfall that stays unsold', bold=True)
put(k2, 'C22', 'Gives 1.0 pt vs. consensus case, about half of row 21')
inp(k2, 'D22', 0.2)
qhead(k2, 24)
put(k2, 'B25', 'Apply to quarter (2027 sailings)')
for c, v in zip(KC, [0, 0, 1, 1, 1, 1]):
    inp(k2, f'{c}25', v, '0')
put(k2, 'B26', 'KPI 2: occupancy vs. consensus case, base (pts)', bold=True)
put(k2, 'B27', 'KPI 2: occupancy vs. consensus case, bear (pts)')
for c in KC:
    fx(k2, f'{c}26', f'=-{c}25*$D$11*$D$22*100', PTS, bold=True)
    fx(k2, f'{c}27', f'=-{c}25*$D$12*$D$22*100', PTS)

sec(k2, 'B29', 'C. Cabins sold late: the close-in discount (price, 2H27 sailings)')
put(k2, 'B30', 'Close-in discount: NCL cut on sailings within ~3 months (4Q26)')
put(k2, 'C30', 'KPI 1 build, 4Q26, NCL brand Caribbean/other mix')
fx(k2, 'D30', f"={K1}E13*{K1}E16+(1-{K1}E13)*{K1}E17")
put(k2, 'B31', 'Apply to quarter (2H27: not yet visible in the fare panel)')
for c, v in zip(KC, [0, 0, 0, 0, 1, 1]):
    inp(k2, f'{c}31', v, '0')
put(k2, 'B32', 'Thesis 2 price effect, base (pts of NCLH ticket revenue per passenger day)', bold=True)
put(k2, 'C32', 'shortfall x share sold late x close-in discount x NCL weight')
put(k2, 'B33', 'Thesis 2 price effect, bear')
for c in KC:
    fx(k2, f'{c}32', f"={c}31*$D$11*(1-$D$22)*$D$30*{K1}$D$33", PCT2, bold=True)
    fx(k2, f'{c}33', f"={c}31*$D$12*(1-$D$22)*$D$30*{K1}$D$33", PCT2)

sec(k2, 'B35', 'D. Best Price Guarantee claims (not in the base case; switch for sensitivity)')
put(k2, 'B36', 'Include in the Net Yield Build? (1 = yes)')
inp(k2, 'D36', cal.get('leak_switch', 0), '0')
put(k2, 'B37', 'Claim rate: booked guests who reprice before final payment')
inp(k2, 'D37', 1 / 3)
put(k2, 'B38', 'Claim rate: paid-up guests who take an upgrade or credit')
inp(k2, 'D38', 1 / 3)
put(k2, 'B39', 'Share of claims taken as future-cruise credit (rest as upgrades)')
inp(k2, 'D39', 0.5)
put(k2, 'B40', 'Credits redeemed evenly over the next quarters (valid two years)')
inp(k2, 'D40', 8, '0')
qhead(k2, 41)
put(k2, 'B42', "NCL brand fare cut on the quarter's sailings")
put(k2, 'B43', 'Share of booked guests already past final payment when cut')
put(k2, 'B44', 'Share of cabins sold before the cuts (base)')
put(k2, 'B45', 'Repricing cost before final payment (pts of NCLH ticket revenue)')
put(k2, 'B46', 'Credits issued to paid-up guests (pts of NCLH ticket revenue)')
put(k2, 'B47', 'Ticket revenue, consensus case ($mm)')
put(k2, 'B48', 'Credits issued ($mm)')
put(k2, 'B49', 'Credits redeemed ($mm)')
put(k2, 'B50', 'Redemption cost (pts of ticket revenue)')
put(k2, 'B51', 'Total claim cost (pts of ticket revenue)', bold=True)
put(k2, 'B52', 'Applied in the Net Yield Build (x switch)')
for i, (c, n) in enumerate(zip(KC, NYB_COLS)):
    fx(k2, f'{c}42', f"={K1}{c}13*{K1}{c}16+(1-{K1}{c}13)*{K1}{c}17")
    inp(k2, f'{c}43', [1, 1, 0.3, 0, 0, 0][i])
    fx(k2, f'{c}44', f"=1-{K1}{c}27")
    fx(k2, f'{c}45', f"={c}44*(1-{c}43)*$D$37*{c}42*{K1}$D$33", PCT2)
    fx(k2, f'{c}46', f"={c}44*{c}43*'Fare Panel'!$F$196*$D$38*$D$39*{c}42*{K1}$D$33", PCT2)
    fx(k2, f'{c}47', f"='Net Yield Build'!{n}32", MM)
    fx(k2, f'{c}48', f'=-{c}46*{c}47', MM)
    fx(k2, f'{c}49', '=0' if i == 0 else f'=SUM($D$48:{KC[i - 1]}48)/$D$40', MM)
    fx(k2, f'{c}50', f'=-{c}49/{c}47', PCT2)
    fx(k2, f'{c}51', f'={c}45+{c}50', PCT2, bold=True)
    fx(k2, f'{c}52', f'={c}51*$D$36', PCT2)
put(k2, 'B53', 'Memo: 2027 EBITDA cost if switched on ($mm)')
fx(k2, 'D53', "=SUMPRODUCT(F51:I51,'Net Yield Build'!N59:Q59)", MM)
put(k2, 'B54', 'Memo: 2027 net yield cost if switched on (pts)')
fx(k2, 'D54', "=D53/'Net Yield Build'!U64", PCT2)

sec(k2, 'B56', 'E. Dated checks on the booking shortfall')
put(k2, 'B57', 'Sept 30, 2026 advance ticket sales: our forecast ($mm, Nov 4 10-Q)')
fx(k2, 'D57', '=Deposits!C70', MM0)
put(k2, 'B58', 'Supports the thesis if below ($mm)')
fx(k2, 'D58', '=Deposits!C72', MM0)
put(k2, 'B59', 'Thesis 2 wrong if at or above ($mm)')
fx(k2, 'D59', '=Deposits!C73', MM0)
put(k2, 'B60', 'Back-test: share of quarters the indicator called the direction of net yield')
fx(k2, 'D60', '=Deposits!F31')
note(k2, 'B97', 'Sources: SEC XBRL (customer deposits); NCLH 8-Ks and 2Q26 call; ncl.com Best Price Guarantee terms; fare panel. '
     'Blue on gray = assumptions.')

sec(k2, 'B64', 'F. Timing: since COVID, deposits lead net yield by about four quarters')
put(k2, 'B65', 'Advance ticket sales as a share of annual revenue, 2016-18 average')
fx(k2, 'D65', "=AVERAGE('Balance (A)'!D26:F26)")
put(k2, 'B66', 'Same, 2023-25 average')
fx(k2, 'D66', "=AVERAGE('Balance (A)'!K26:M26)")
note(k2, 'B67', 'Guests book further ahead than before COVID, so deposits run further ahead of sailings, and of net yield.')
for c, h in zip('BCDEFG', ['Quarter end', 'Deposits per future berth-day y/y', 'Net yield y/y, 1 qtr later', '2 qtrs later',
                           '3 qtrs later', '4 qtrs later']):
    put(k2, f'{c}69', h, bold=True, wrap=True)
for c, h in zip('IJK', ['Net yield series', 'Net yield y/y', 'Source']):
    put(k2, f'{c}69', h, bold=True, wrap=True)
k2.row_dimensions[69].height = 28
k2.column_dimensions['J'].width = 11
k2.column_dimensions['K'].width = 14
qlab = ['Mar-24', 'Jun-24', 'Sep-24', 'Dec-24', 'Mar-25', 'Jun-25', 'Sep-25', 'Dec-25', 'Mar-26', 'Jun-26', 'Sep-26', 'Dec-26',
        'Mar-27', 'Jun-27']
for i, q in enumerate(qlab):
    r = 70 + i
    put(k2, f'I{r}', q)
    if i < 10:
        fx(k2, f'J{r}', f"='Multiple vs Net Yield'!I{18 + i}")
        put(k2, f'K{r}', 'reported')
    elif i == 10:
        fx(k2, f'J{r}', '=Consensus!C41')
        put(k2, f'K{r}', '3Q26 guide')
    elif i == 11:
        fx(k2, f'J{r}', '=Consensus!C42')
        put(k2, f'K{r}', '4Q26 implied')
    else:
        fx(k2, f'J{r}', '=""', 'General')
        put(k2, f'K{r}', 'not yet known')
for i in range(10):
    r = 70 + i
    fx(k2, f'B{r}', f'=Deposits!B{17 + i}', 'mmm-yy')
    fx(k2, f'C{r}', f'=Deposits!F{17 + i}')
    for k, c in enumerate('DEFG', start=1):
        fx(k2, f'{c}{r}', f'=IF(INDEX($J$70:$J$83,{i + 1 + k})="","",INDEX($J$70:$J$83,{i + 1 + k}))')
put(k2, 'B85', 'Lead (quarters)', bold=True)
for c, h in zip('DEFG', ['1', '2', '3', '4']):
    put(k2, f'{c}85', h, bold=True).alignment = Alignment(horizontal='right')
put(k2, 'B86', 'Quarters with a known outcome')
put(k2, 'B87', 'Direction called correctly')
put(k2, 'B88', 'Correlation')
for c in 'DEFG':
    fx(k2, f'{c}86', f'=COUNT({c}70:{c}79)', '0')
    fx(k2, f'{c}87', f'=SUMPRODUCT(ISNUMBER({c}70:{c}79)*(({c}70:{c}79>0)=($C$70:$C$79>0)))', '0')
    fx(k2, f'{c}88', f'=CORREL($C$70:$C$79,{c}70:{c}79)', '0.00')
put(k2, 'B89', 'Direction called at a 4-quarter lead, reported outcomes only (of 6)')
fx(k2, 'G89', '=SUMPRODUCT(--((C70:C75>0)=(G70:G75>0)))', '0')
note(k2, 'B90', 'Outcomes include the 3Q26 guide and 4Q26 implied net yield. 2016-19 is not informative: the indicator and net yield were '
     'positive in every quarter.')
put(k2, 'B92', 'Fit at a 4-quarter lead (net yield = intercept + slope x indicator)', bold=True)
for c, h in zip('DEF', ['Indicator fit', 'Our base', 'Consensus case']):
    put(k2, f'{c}92', h, bold=True).alignment = Alignment(horizontal='right')
put(k2, 'B93', 'Slope / intercept')
fx(k2, 'D93', '=SLOPE(G70:G79,C70:C79)', '0.00')
fx(k2, 'E93', '=INTERCEPT(G70:G79,C70:C79)', PCT2)
put(k2, 'B94', 'Implied 1Q27 net yield (Mar 31, 2026 reading) / 2Q27 (Jun 30 reading)')
fx(k2, 'D94', '=E93+D93*C78', PCT)
fx(k2, 'E94', '=E93+D93*C79', PCT)
put(k2, 'B95', '1H27 net yield: indicator fit vs. our base vs. consensus case', bold=True)
fx(k2, 'D95', '=AVERAGE(D94:E94)', PCT, bold=True)
fx(k2, 'E95', "=AVERAGE('Net Yield Build'!N58:O58)", PCT, bold=True)
fx(k2, 'F95', "=AVERAGE('Net Yield Build'!N31:O31)", PCT, bold=True)
note(k2, 'B96', 'Eight quarters, two of them guidance: a cross-check on direction and size, not a forecast input.')

# ---------------------------------------------------------------- 5. Net Yield Build rewiring
ny = wb['Net Yield Build']
put(ny, 'B6', 'Capacity days growth y/y (3Q-4Q26 guidance; 2027 Fleet Qtr Build)')
for n, fc in zip(['N', 'O', 'P', 'Q'], ['L', 'M', 'N', 'O']):
    fx(ny, f'{n}6', f"='Fleet Qtr Build'!{fc}12")
for r, lab, srow in [(7, 'Caribbean share of NCLH capacity (KPI 1 build)', 10),
                     (8, 'Observed NCL fare change: Caribbean (KPI 1 build)', 16),
                     (9, 'Observed NCL fare change: other regions (KPI 1 build)', 17),
                     (10, 'Full-fare commission step-up (pts of ticket; KPI 1 build)', 49)]:
    put(ny, f'B{r}', lab)
    for n, c in zip(NYB_COLS, KC):
        fx(ny, f'{n}{r}', f'={K1}{c}{srow}')

blocks = {  # case: (first driver row, occupancy src, T1 src, T2 src, leak?)
    'bull': (13, None, f'{K1}{{c}}40', None, False),
    'base': (40, f'{K2}{{c}}26', f'{K1}{{c}}39', f'{K2}{{c}}32', True),
    'bear': (67, f'{K2}{{c}}27', f'{K1}{{c}}41', f'{K2}{{c}}33', True),
}
calib = cal.get('m_underlying', {})
for case, (r0, occ, t1src, t2src, leak) in blocks.items():
    ro, ru, rt1, rmix, rt2, rob = r0, r0 + 1, r0 + 2, r0 + 3, r0 + 4, r0 + 5
    calc0 = r0 + 9                                           # first like-for-like row (22 / 49 / 76)
    if occ:
        put(ny, f'B{ro}', 'Occupancy change y/y (pts): consensus case + KPI 2 (2027)')
        for n, c in zip(NYB_COLS[2:], KC[2:]):
            fx(ny, f'{n}{ro}', f'={n}13+{occ.format(c=c)}', PTS)
    if case == 'base':
        put(ny, f'B{ru}', 'Underlying price on sold cabins (2H26 calibrated to guidance; 2027 = consensus case)')
        for n in NYB_COLS[2:]:
            fx(ny, f'{n}{ru}', f'={n}14')
    if case in calib:
        inp(ny, f'M{ru}', calib[case])
    put(ny, f'B{rt1}', 'Thesis 1: fare cuts on unsold cabins (KPI 1 build)')
    for n, c in zip(NYB_COLS, KC):
        fx(ny, f'{n}{rt1}', '=' + t1src.format(c=c))
    put(ny, f'B{rt2}', 'Thesis 2: late discounts on the booking shortfall (KPI 2 build)')
    for n, c in zip(NYB_COLS, KC):
        if t2src:
            fx(ny, f'{n}{rt2}', '=' + t2src.format(c=c))
        else:
            inp(ny, f'{n}{rt2}', 0)
    put(ny, f'B{calc0}', 'Underlying price change')
    put(ny, f'B{calc0 + 1}', 'Thesis effects (Thesis 1 + Thesis 2)')
    put(ny, f'B{calc0 + 2}', 'KPI 1: like-for-like ticket price change y/y', bold=True)
    for n in NYB_COLS:
        fx(ny, f'{n}{calc0}', f'={n}{ru}')
        fx(ny, f'{n}{calc0 + 1}', f'={n}{rt1}+{n}{rt2}')
        fx(ny, f'{n}{calc0 + 2}', f'={n}{calc0}+{n}{calc0 + 1}', bold=True)
    tick = calc0 + 3
    prior = {'L': 'H', 'M': 'I', 'N': 'J', 'O': 'K', 'P': 'L', 'Q': 'M'}
    for n, c in zip(NYB_COLS, KC):
        lk = f"+{K2}{c}52" if leak else ''
        fx(ny, f'{n}{tick}', f'={prior[n]}{tick}*(1+{n}{calc0 + 2}+{n}{rmix}{lk})', DOL)

# ---------------------------------------------------------------- 6. Bridge Build (live consensus -> base bridge)
bb = new_sheet('Bridge Build', {'A': 2, 'B': 58, 'C': 30, 'D': 11, 'E': 11, 'F': 11, 'G': 11, 'H': 11, 'I': 11, 'J': 12})
bar(bb, 'B2', 'Bridge Build: 2027 EBITDA from the consensus case to our base case, one KPI at a time')
note(bb, 'B3', 'Each block reruns the quarterly net yield with one more group of base-case inputs switched on, in order. The change '
     'between blocks is the EBITDA attributable to that group. The last step must equal DCF!E42 (checks in row 13).')
steps = [('Consensus case (bull inputs)', {}),
         ('2H26 exit (3Q-4Q26 inputs)', {'X'}),
         ('Thesis 1 -> KPI 1: fare cuts on unsold cabins', {'X', 'T1'}),
         ('Thesis 2 -> KPI 1: late discounts on the shortfall', {'X', 'T1', 'T2'}),
         ('Thesis 2 -> KPI 2: occupancy', {'X', 'T1', 'T2', 'OCC'}),
         ('Other: onboard spend growth', {'X', 'T1', 'T2', 'OCC', 'OB'})]
put(bb, 'B5', 'Step', bold=True)
put(bb, 'D5', '2027 EBITDA', bold=True)
put(bb, 'E5', 'Change', bold=True)
B0 = 14                     # first block row
BLK = 9                     # rows per block
NCC_C = "'Operating Model'!$G$45*(1+'Operating Model'!$H$11)*(1+'Operating Model'!$I$11)*{cap}"
FUEL_C = "'Operating Model'!$H$47*({cap}/'Net Yield Build'!$T$19)*0.99*'Operating Model'!$I$16/1000"
NCC_B = "'Operating Model'!$G$45*(1+'Operating Model'!$H$12)*(1+'Operating Model'!$I$12)*{cap}"
FUEL_B = "'Operating Model'!$H$47*({cap}/'Net Yield Build'!$T$19)*0.99*'Operating Model'!$I$17/1000"
prior_actual = {'D': 'H', 'E': 'I', 'F': 'J', 'G': 'K'}
for s, (lab, on) in enumerate(steps):
    top = B0 + s * BLK
    sec(bb, f'B{top}', f'Block {s}: {lab}')
    qhead(bb, top, extra={'J': 'FY2027E'})
    rows = dict(occd=top + 1, occ=top + 2, lfl=top + 3, tick=top + 4, onb=top + 5, gm=top + 6, ebitda=top + 7)
    put(bb, f'B{rows["occd"]}', 'Occupancy change y/y (pts)')
    put(bb, f'B{rows["occ"]}', 'Occupancy')
    put(bb, f'B{rows["lfl"]}', 'Like-for-like ticket price change + mix')
    put(bb, f'B{rows["tick"]}', 'Ticket revenue per passenger day ($)')
    put(bb, f'B{rows["onb"]}', 'Onboard revenue per passenger day ($)')
    put(bb, f'B{rows["gm"]}', 'Adjusted gross margin ($mm)')
    put(bb, f'B{rows["ebitda"]}', 'Adjusted EBITDA ($mm)', bold=True)
    for i, (c, n) in enumerate(zip(KC, NYB_COLS)):
        late = i < 2
        sw = lambda g: ('X' in on) if late else (g in on)
        occ_src = f"'Net Yield Build'!{n}{40 if sw('OCC') else 13}"
        und = f"'Net Yield Build'!{n}{41 if sw('X') else 14}"
        t1 = f"'Net Yield Build'!{n}{42 if sw('T1') else 15}"
        t2 = f"('Net Yield Build'!{n}44+{K2}{c}52)" if sw('T2') else f"'Net Yield Build'!{n}17"
        obg = f"'Net Yield Build'!{n}{45 if sw('OB') else 18}"
        pr = (lambda r: f"'Net Yield Build'!{prior_actual[c]}{r}") if c in prior_actual else \
             (lambda r, c=c: f'{KC[KC.index(c) - 2]}{r}')
        fx(bb, f'{c}{rows["occd"]}', f'={occ_src}', PTS)
        prior_occ = f"'Net Yield Build'!{prior_actual[c]}20" if c in prior_actual else f'{KC[i - 4]}{rows["occ"]}'
        fx(bb, f'{c}{rows["occ"]}', f'={prior_occ}+{c}{rows["occd"]}/100', PCT)
        fx(bb, f'{c}{rows["lfl"]}', f"={und}+{t1}+{t2}+'Net Yield Build'!{n}16", PCT2)
        prior_t = f"'Net Yield Build'!{prior_actual[c]}25" if c in prior_actual else f'{KC[i - 4]}{rows["tick"]}'
        prior_o = f"'Net Yield Build'!{prior_actual[c]}26" if c in prior_actual else f'{KC[i - 4]}{rows["onb"]}'
        fx(bb, f'{c}{rows["tick"]}', f'={prior_t}*(1+{c}{rows["lfl"]})', DOL)
        fx(bb, f'{c}{rows["onb"]}', f'={prior_o}*(1+{obg})', DOL)
        fx(bb, f'{c}{rows["gm"]}', f"='Net Yield Build'!{n}19*{c}{rows['occ']}*({c}{rows['tick']}*(1-'Net Yield Build'!{n}54)"
                                    f"+{c}{rows['onb']}*(1-'Net Yield Build'!{n}55))", MM)
    fx(bb, f'J{rows["gm"]}', f'=SUM(F{rows["gm"]}:I{rows["gm"]})', MM)
    cap = "'Net Yield Build'!$U$19"
    fx(bb, f'J{rows["ebitda"]}', f'=J{rows["gm"]}-{NCC_C.format(cap=cap)}-{FUEL_C.format(cap=cap)}', MM, bold=True)
    put(bb, f'B{6 + s}', lab)
    fx(bb, f'D{6 + s}', f'=J{rows["ebitda"]}', MM)
    if s:
        fx(bb, f'E{6 + s}', f'=D{6 + s}-D{5 + s}', MM)
last = B0 + 5 * BLK
put(bb, 'B12', 'Other: unit costs and fuel')
fx(bb, 'D12', f"=J{last + 6}-{NCC_B.format(cap=chr(39) + 'Net Yield Build' + chr(39) + '!$U$19')}"
              f"-{FUEL_B.format(cap=chr(39) + 'Net Yield Build' + chr(39) + '!$U$19')}", MM)
fx(bb, 'E12', '=D12-D11', MM)
put(bb, 'B13', 'Checks: block 0 = DCF!E41, final = DCF!E42 (both should be 0)', bold=True)
fx(bb, 'D13', '=D6-DCF!E41', MM)
fx(bb, 'E13', '=D12-DCF!E42', MM)

# Model vs Consensus bridge now links to the live build, ordered: consensus EPS -> Street EBITDA -> ours
mc = wb['Model vs Consensus']
for r in range(22, 42):
    for c in 'BCDEF':
        mc[f'{c}{r}'].value = None
sec(mc, 'B22', '2027 Adjusted EBITDA bridge: consensus EPS case -> consensus case at run-rate costs (Street EBITDA) -> base case')
put(mc, 'B23', 'Step', bold=True)
put(mc, 'C23', '$mm', bold=True)
put(mc, 'D23', 'Street-implied', bold=True)
put(mc, 'E23', 'Difference', bold=True)
put(mc, 'B24', 'Consensus EPS case (bull inputs: Street revenue, unit costs -0.5%, fuel $760)')
fx(mc, 'C24', "='Bridge Build'!D6", MM)
put(mc, 'B25', 'Run-rate unit costs and fuel (+1.0%, $800; same in our case)')
fx(mc, 'C25', "='Bridge Build'!E12", MM)
put(mc, 'B26', 'Consensus case at run-rate costs (compare: Street-implied EBITDA)', bold=True)
fx(mc, 'C26', '=C24+C25', MM, bold=True)
fx(mc, 'D26', '=Consensus!C57', MM)
fx(mc, 'E26', '=C26/D26-1', PCT)
steps_mc = [('Fare panel -> 4Q26 exit rate (Q4 2026 sailings cut 36%)', 7),
            ('Fare panel -> KPI 1: 1H27 ticket price (Thesis 1)', 8),
            ('Deposits -> KPI 1: 2H27 late discounts on the shortfall (Thesis 2)', 9),
            ('Deposits -> KPI 2: occupancy (Thesis 2)', 10),
            ('Guest mix from discount-filled ships -> onboard spend (both theses)', 11)]
for i, (lab, br) in enumerate(steps_mc):
    put(mc, f'B{27 + i}', lab)
    fx(mc, f'C{27 + i}', f"='Bridge Build'!E{br}", MM)
put(mc, 'B32', 'Model base case', bold=True)
fx(mc, 'C32', '=DCF!E42', MM, bold=True)
put(mc, 'B33', 'Check: steps sum to the difference')
fx(mc, 'C33', '=C26+SUM(C27:C31)-C32', MM)
put(mc, 'B34', 'Calibration: consensus case 2027 revenue vs. Street revenue ($mm)')
fx(mc, 'C34', "='Net Yield Build'!U34", MM)
fx(mc, 'D34', '=Consensus!D6', MM)
fx(mc, 'E34', '=C34/D34-1', PCT)
put(mc, 'B35', 'Consensus EPS at Street EBITDA and the guided $870M interest')
fx(mc, 'C35', '=Consensus!C58', DOL)
note(mc, 'B36', 'Costs are separable from the revenue steps, so the order of the bridge does not change any step. Bridge Build reruns the '
     'quarterly net yield one input group at a time; its checks (row 13) tie to DCF!E41 and DCF!E42.')

# ---------------------------------------------------------------- 7. Thesis to KPIs (front map)
tk = new_sheet('Thesis to KPIs', {'A': 2, 'B': 58, 'C': 26, 'D': 24, 'E': 18, 'F': 14, 'G': 14, 'H': 14, 'I': 13}, index=1)
NY = "'Net Yield Build'!"
bar(tk, 'B2', 'Thesis to KPIs: what the alternative data shows, which model inputs it moves, and where we differ from consensus', 'I')
note(tk, 'B3', 'Net yield = occupancy x net revenue per passenger day. Two data sets drive the view: the fare panel moves ticket price '
     '(KPI 1) and deposits move occupancy (KPI 2) and late discounts. Unit costs are run-rate in both our case and the Street\'s EBITDA.')

sec(tk, 'B5', 'A. FY2027: consensus case vs. our base case')
for c, h in zip('CDE', ['Consensus case', 'Base case', 'Difference']):
    put(tk, f'{c}5', h, bold=True).alignment = Alignment(horizontal='right')
fy = [('Occupancy', 'U20', 'U47', PCT, 'pts'),
      ('KPI 1 input: ticket revenue per passenger day ($)', 'U25', 'U52', DOL, '%'),
      ('Onboard revenue per passenger day ($)', 'U26', 'U53', DOL, '%'),
      ('Net revenue per passenger day ($)', 'U29', 'U56', DOL, '%'),
      ('Net yield ($ per capacity day)', 'U30', 'U57', DOL, '%'),
      ('Net yield growth y/y', 'U31', 'U58', PCT, 'pts')]
for i, (lab, a, b, fmt, kind) in enumerate(fy):
    r = 6 + i
    put(tk, f'B{r}', lab, bold=(i in (0, 1)))
    fx(tk, f'C{r}', f'={NY}{a}', fmt)
    fx(tk, f'D{r}', f'={NY}{b}', fmt)
    fx(tk, f'E{r}', f'=(D{r}-C{r})*100' if kind == 'pts' else f'=D{r}/C{r}-1', PTS if kind == 'pts' else PCT)
put(tk, 'B12', 'Adjusted EBITDA ($mm)', bold=True)
fx(tk, 'C12', '=DCF!E41', MM0)
fx(tk, 'D12', '=DCF!E42', MM0)
fx(tk, 'E12', '=D12/C12-1', PCT)

sec(tk, 'B14', 'B. The two KPIs by quarter')
qhead(tk, 14, cols=['C', 'D', 'E', 'F', 'G', 'H'])
kq = [('KPI 1: like-for-like ticket price y/y, consensus case', 24, PCT),
      ('KPI 1: ours', 51, PCT),
      ('   of which Thesis 1 (fare cuts on unsold cabins)', 42, PCT2),
      ('   of which Thesis 2 (late discounts on the shortfall)', 44, PCT2),
      ('KPI 2: occupancy change y/y (pts), consensus case', 13, PTS),
      ('KPI 2: ours', 40, PTS),
      ('Net yield y/y, consensus case', 31, PCT),
      ('Net yield y/y, ours', 58, PCT)]
for i, (lab, row, fmt) in enumerate(kq):
    r = 15 + i
    put(tk, f'B{r}', lab, bold=lab.startswith('KPI') and 'ours' in lab)
    for c, n in zip('CDEFGH', NYB_COLS):
        fx(tk, f'{c}{r}', f'={NY}{n}{row}', fmt)

sec(tk, 'B24', 'C. From the data to the model: the 2027 EBITDA bridge from the consensus case at run-rate costs')
heads = ['Step', 'What the data shows', 'Model input it sets', 'KPI', 'Consensus case', 'Ours', '2027 EBITDA ($mm)']
cols = ['B', 'C', 'D', 'E', 'F', 'G', 'H']
for c, h in zip(cols, heads):
    put(tk, f'{c}25', h, bold=True, wrap=True)
tk.row_dimensions[25].height = 28
MC = "'Model vs Consensus'!"
t_rows = [
    ('Consensus case at run-rate costs', f'="Reproduces Street revenue ("&TEXT({MC}E34,"+0.0%;-0.0%")&") and EBITDA ("&TEXT({MC}E26,"+0.0%;-0.0%")&")"',
     'Bull revenue drivers with our unit costs and fuel', '', '', '', f'={MC}C26'),
    ('Fare panel: 4Q26 exit', f'="NCL Q4 2026 Caribbean sailings "&TEXT({K1}E16,"0%")&"; 77% of 7+ nights cut 20%+"',
     '4Q26 net yield below the implied guide', 'Net yield 4Q26', f'={NY}M31', f'={NY}M58', f'={MC}C27'),
    ('Fare panel: Thesis 1', f'="1H27 Caribbean fares "&TEXT({K1}F16,"0.0%")&" vs. peers "&TEXT({K1}F19,"0.0%")&"; luxury brands 0%"',
     'Cuts applied to unsold cabins, NCL brand at 65% of ticket revenue', 'KPI 1: 1H27 ticket price',
     f'=AVERAGE({NY}N24:O24)', f'=AVERAGE({NY}N51:O51)', f'={MC}C28'),
    ('Deposits: Thesis 2 (price)', f'="Deposits per future berth-day "&TEXT(Deposits!F26,"0.0%")&"; lead net yield ~4 qtrs since COVID ("&{K2}G87&" of "&{K2}G86&")"',
     '5-pt booking shortfall; 80% sold late at the close-in discount', 'KPI 1: 2H27 ticket price',
     f'=AVERAGE({NY}P24:Q24)', f'=AVERAGE({NY}P51:Q51)', f'={MC}C29'),
    ('Deposits: Thesis 2 (occupancy)', f'="2026: occupancy fell ~"&TEXT({K2}D19,"0.00")&" pts per 1% of indicator decline"',
     '20% of the shortfall sails empty (half the 2026 relationship)', 'KPI 2: 2027 occupancy', f'={NY}U20', f'={NY}U47', f'={MC}C30'),
    ('Both: guest mix', f'="1Q26: occupancy +2.3 pts on price cuts, onboard per guest "&TEXT(\'Quarterly (A)\'!I35/\'Quarterly (A)\'!E35-1,"0.0%")',
     'Onboard spend growth 2% (2025 -0.3%, 1H26 ~+1%)', 'Onboard per guest y/y', f'={NY}N18', f'={NY}N45', f'={MC}C31'),
    ('Our base case', '', '', '', f'={MC}C26', f'={MC}C32', f'={MC}C32'),
]
for i, row in enumerate(t_rows):
    r = 26 + i
    for c, v in zip(cols, row):
        fmt = PCT if c in 'FG' else MM
        if i in (0, 6) and c in 'FG':
            fmt = MM
        if isinstance(v, str) and v.startswith('='):
            fx(tk, f'{c}{r}', v, fmt, bold=i in (0, 6))
            tk[f'{c}{r}'].alignment = Alignment(wrap_text=True, vertical='top')
        elif v:
            put(tk, f'{c}{r}', v, wrap=True, bold=(c == 'B'))
    tk.row_dimensions[r].height = 44

sec(tk, 'B34', 'Consensus EPS vs. Street EBITDA (not part of the theses)')
put(tk, 'B35', 'Unit costs ex fuel per capacity day, 2027 y/y: consensus EPS case / ours')
fx(tk, 'F35', "='Operating Model'!I11")
fx(tk, 'G35', "='Operating Model'!I12")
fx(tk, 'H35', f'={MC}C25', MM)
put(tk, 'B36', 'Fuel per ton, net of hedges: consensus EPS case / ours')
fx(tk, 'F36', "='Operating Model'!I16", '$#,##0')
fx(tk, 'G36', "='Operating Model'!I17", '$#,##0')
put(tk, 'B37', 'Consensus EPS at Street EBITDA and the guided interest')
fx(tk, 'G37', f'={MC}C35', DOL)
note(tk, 'B38', 'Consensus EPS ($1.46-1.67) needs unit costs to fall on top of the Street\'s EBITDA; the Street\'s own EBITDA matches run-rate costs.')

sec(tk, 'B40', 'D. What each KPI is worth, and valuation checks (base case, 2027)')
put(tk, 'B41', '1 pt of ticket price per passenger day: EBITDA ($mm) / per share')
fx(tk, 'C41', f"={NY}U59*(1-{NY}U54)*0.01", MM)
fx(tk, 'D41', '=C41*$C$43', DOL)
put(tk, 'B42', '1 pt of occupancy: EBITDA ($mm) / per share')
fx(tk, 'C42', f"={NY}U46*0.01*{NY}U56", MM)
fx(tk, 'D42', '=C42*$C$43', DOL)
put(tk, 'B43', "Value per share of $1M of 2027 EBITDA (today's multiple)")
fx(tk, 'C43', '=DCF!F25/DCF!L33', '$0.000')
put(tk, 'B44', 'Value if unit costs fall 0.5% as consensus EPS implies')
fx(tk, 'C44', f"=((DCF!E42-{MC}C25)*DCF!F25-DCF!F22)/DCF!L33", DOL)
fx(tk, 'D44', '=C44/DCF!I2-1', PCT)
put(tk, 'B45', 'Value at 9.0x on our EBITDA (multiple risk)')
fx(tk, 'C45', '=(DCF!E42*9-DCF!F22)/DCF!L33', DOL)
fx(tk, 'D45', '=C45/DCF!I2-1', PCT)
put(tk, 'B46', 'EV / TTM EBITDA: Sep 2025 -> Mar 2026 (net yield turned negative)')
fx(tk, 'C46', "='Multiple vs Net Yield'!H24", '0.0x')
fx(tk, 'D46', "='Multiple vs Net Yield'!H26", '0.0x')

sec(tk, 'B48', 'E. Dated checks: when each KPI gets tested')
checks = [('Nov 4, 2026: 3Q26 10-Q', 'Sept 30 advance ticket sales vs. our $2.94B (KPI 2 shortfall)',
           f'="Below "&TEXT({K2}D58,"$#,##0")&"M supports; at or above "&TEXT({K2}D59,"$#,##0")&"M we cover"'),
          ('Late Feb 2027: first 2027 guide', 'Net yield guide vs. the consensus case (KPI 1 and KPI 2)',
           f'="At or below flat supports; consensus case "&TEXT({NY}U31,"+0.0%")'),
          ('Monthly: fare panel re-run', '1H27 Caribbean cuts on unsold cabins (KPI 1)', 'Cuts back near peers would undo Thesis 1'),
          ('May 2027: 1Q27 results', 'Reported ticket revenue per passenger day and occupancy',
           f'="Ours: occupancy "&TEXT({NY}N47,"0.0%")&", ticket price "&TEXT({NY}N51,"0.0%")')]
for i, (d, what, test) in enumerate(checks):
    r = 49 + i
    put(tk, f'B{r}', d)
    put(tk, f'C{r}', what)
    if test.startswith('='):
        fx(tk, f'G{r}', test, 'General')
    else:
        put(tk, f'G{r}', test)

sec(tk, 'B54', 'F. Critiques we expect, and where the model answers them')
put(tk, 'B55', 'Critique', bold=True)
put(tk, 'C55', 'Answer', bold=True)
crit = [
    ('Lowest fares are not realized prices.', 'Realized ticket revenue per passenger day already fell 5.6% in 1H26 (KPI 2 Booking Build D7). Cuts are applied only to unsold cabins (KPI 1 Build row 27).'),
    ('It is the war, or NCL simply moved capacity to the Caribbean.', 'Royal and Carnival also added Caribbean ships and held price (-1%, 0%); they are cutting Europe instead (Fare Panel A).'),
    ('Oceania and Regent are a third of the revenue and are not cutting.', 'Cuts are weighted by brand: NCL at 65% of ticket revenue, calibrated to reported revenue at 66% (KPI 1 Build section H).'),
    ('The deposit decline is just a shift to shorter cruises.', 'Short (3-5 night) sailings are about a third of 4Q26 Caribbean itineraries at NCL, Royal and Carnival alike; NCL cuts both lengths (Fare Panel C).'),
    ('Deposits turned a year early, so the signal failed.', 'Since COVID the lead is about four quarters (7 of 8 at a 4-quarter lead vs. 6 of 10 at one); KPI 2 Build section F.'),
    ('The theses do not explain the gap to consensus.', 'At run-rate costs the consensus case matches Street revenue and EBITDA; the full $208M gap from there is fare-panel and deposit driven (section C).'),
    ('The shortfall size and empty-cabin share are made up.', 'Shortfall = midpoint of the deposit-implied 3-8 pt range; empty share is half the 2026 occupancy relationship (KPI 2 Build A-B).'),
    ('The multiple could expand as the cycle bottoms.', 'At 9.0x our EBITDA is worth less than today\'s price; NCLH\'s multiple fell from 10.0x to 8.7x as net yield turned negative (section D).'),
    ('Cost savings could beat the plan.', 'Shown: if unit costs fall 0.5%, value rises to the figure in D44, still below the share price.'),
    ('Crowded short (20.7% of float).', 'Sizing: a third to start, add on oil-driven rallies, puts for part, stop above $19 (memo).'),
]
for i, (q, a) in enumerate(crit):
    r = 56 + i
    put(tk, f'B{r}', q, wrap=True)
    put(tk, f'C{r}', a, wrap=True)
    tk.merge_cells(f'C{r}:I{r}')
    tk.row_dimensions[r].height = 28
note(tk, 'B67', 'Green = links; black = formulas. Every number on this sheet comes from the KPI builds, the Net Yield Build or the DCF.')

dcf = wb['DCF']
fx(dcf, 'F22', "='Operating Model'!H88", MM0)
put(dcf, 'B22', 'Net debt at YE2026E used in valuation (= model)')

# ---------------------------------------------------------------- 8. front page, exhibits, sheet order
mm = wb['Model>>>']
put(mm, 'C21', 'Price-implied steady-state EBITDA at 9% ($mm)')
fx(mm, 'I21', '=DCF!E31', MM0)
sheet_notes = [
    ('Thesis to KPIs', 'Start here: data -> KPI -> model -> consensus, with dated checks'),
    ('DCF', 'Valuation: case switch (C9), WACC, reverse DCF, EV/EBITDA, DCF cross-check, scenarios, sensitivity'),
    ('Model vs Consensus / Bridge Build', 'Model vs. consensus; live 2027 EBITDA bridge, one step per KPI'),
    ('Operating Model', 'Annual model FY2024A-FY2030E: capacity, net yield, revenue, costs, P&L, cash flow, ROIC'),
    ('Net Yield Build', 'Quarterly net yield for the three cases; KPI rows 24 / 51 / 78 (price) and 13 / 40 / 67 (occupancy)'),
    ('KPI 1 Price Build', 'Thesis 1: fare cuts on unsold cabins by brand and region -> ticket price per passenger day'),
    ('KPI 2 Booking Build', 'Thesis 2: booking shortfall from deposits -> occupancy and late discounts; price-guarantee claims'),
    ('Fleet Build / Fleet Qtr Build', 'Ship-level berths; ship x quarter capacity days (504 rows) for quarterly capacity growth'),
    ('Consensus', 'Annual and quarterly consensus, ratings, targets, guidance, decoding the target'),
    ('Quarterly (A) / Income / Cash Flow / Balance', 'Historical results: 8-K releases and SEC XBRL'),
    ('Fare Panel / Deposits', 'Raw fare panel (4,783 itineraries, summarized) and the deposit indicator (back-test, peers, Nov 4 forecast)'),
    ('Factor & Positioning / Multiple vs Net Yield', 'Return decomposition, news days, oil test, short interest; multiple vs. net yield'),
    ('Exhibit Data', 'Data behind the memo charts'),
]
legend = [mm[f'C{r}'].value for r in range(25, 30)]
for r in range(25, 60):
    for c in 'CDEFGHI':
        mm[f'{c}{r}'].value = None
put(mm, 'C27', legend[0], bold=True, size=11)
for i, (txt, col) in enumerate(zip(legend[1:], [BLUE, BLUE, BLACK, GREEN])):
    put(mm, f'C{28 + i}', txt, color=col, size=11, fill=GRAY if i == 1 else None)
put(mm, 'C24', 'KPI 1: 1H27 ticket price y/y, consensus case / ours')
fx(mm, 'I24', "=AVERAGE('Net Yield Build'!N51:O51)", PCT)
fx(mm, 'H24', "=AVERAGE('Net Yield Build'!N24:O24)", PCT)
put(mm, 'C25', 'KPI 2: 2027 occupancy, consensus case / ours')
fx(mm, 'I25', "='Net Yield Build'!U47", PCT)
fx(mm, 'H25', "='Net Yield Build'!U20", PCT)
put(mm, 'C33', 'Sheets', bold=True, size=11)
for i, (a, b) in enumerate(sheet_notes):
    put(mm, f'C{34 + i}', a, size=11)
    put(mm, f'E{34 + i}', b, size=11)

ex = wb['Exhibit Data']
sec(ex, 'B36', 'Exhibit: KPIs by quarter, consensus case vs. ours (memo chart)')
qhead(ex, 37, cols=['C', 'D', 'E', 'F', 'G', 'H'])
for i, (lab, row, fmt) in enumerate([('Ticket price y/y, consensus case', 24, PCT), ('Ticket price y/y, ours', 51, PCT),
                                      ('Occupancy y/y (pts), consensus case', 13, PTS), ('Occupancy y/y (pts), ours', 40, PTS)]):
    put(ex, f'B{38 + i}', lab)
    for c, n in zip('CDEFGH', NYB_COLS):
        fx(ex, f'{c}{38 + i}', f'={NY}{n}{row}', fmt)
sec(ex, 'B43', 'Exhibit: price-implied EBITDA ($mm)')
for i, (lab, f) in enumerate([('2025 actual', "='Operating Model'!G52"), ('March 2026 guide (2026)', '=Consensus!C44'),
                              ('Price-implied at 9%', '=DCF!E31'), ('Average target needs', '=Consensus!C67'),
                              ('Street 2027', '=Consensus!C57'), ('Ours 2027', "='Operating Model'!I52")]):
    put(ex, f'B{44 + i}', lab)
    fx(ex, f'C{44 + i}', f, MM0)

del wb['Alt Data']
order = ['Model>>>', 'Thesis to KPIs', 'DCF', 'Model vs Consensus', 'Bridge Build', 'Operating Model', 'Net Yield Build',
         'KPI 1 Price Build', 'KPI 2 Booking Build', 'Fleet Build', 'Fleet Qtr Build', 'Consensus', 'Filings>>>', 'Quarterly (A)',
         'Income (A)', 'Cash Flow (A)', 'Balance (A)', 'Alt Data>>>', 'Fare Panel', 'Deposits', 'Factor & Positioning',
         'Multiple vs Net Yield', 'Exhibit Data']
wb._sheets = [wb[n] for n in order]
wb.active = 0
for ws in wb.worksheets:
    ws.sheet_view.tabSelected = ws.title == 'Model>>>'

# no stale references to removed or renamed sheets
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and ("Alt Data Build'" in c.value or "'Alt Data'!" in c.value):
                raise SystemExit(f'stale reference {ws.title}!{c.coordinate}: {c.value}')
wb.save(out)
print('saved', out)
